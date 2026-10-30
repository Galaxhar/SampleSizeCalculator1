"""Create the reviewed protocol and client note, preserving the uploaded original.

Changes: recruitment totals, sample-size section, central R/Excel randomization,
and cited references. All inserted/replaced protocol text is highlighted yellow.
Uses lxml and python-docx; these are editorial, not statistical dependencies.
"""
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from lxml import etree as ET
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'references/original/projeto.docx'
TARGET = ROOT / 'docs/projeto_revisado_calculo_randomizacao.docx'
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W}
Q = lambda name: f'{{{W}}}{name}'
SAMPLE = [
    'O tamanho amostral foi calculado para a comparação do IAH ao final da intervenção (T1; semana 4) por análise de covariância (ANCOVA), ajustada pelo IAH basal e pelo sexo, com alfa bilateral de 5% e poder estatístico de 80% (Vickers e Altman, 2001; Borm et al., 2007). Adotou-se desvio-padrão final de 8,7 eventos/h, derivado da raiz da média das variâncias dos grupos acupuntura e sham de Freire et al. (2007), cujos desvios-padrão finais foram 5,6 e 11,0 eventos/h. A correlação basal–pós de 0,5 foi assumida como hipótese de planejamento, com apoio indireto em uma correlação de aproximadamente 0,56, derivada dos desvios-padrão basal, final e da mudança de um ensaio em AOS moderada, realizado com outro tratamento e seguimento de quatro meses (Puhan et al., 2006); essa estimativa não confirma a correlação para o presente ensaio. A diferença-alvo de 6 eventos/h representa a estimativa clínica de redução adicional com eletroacupuntura em relação ao sham, contextualizada por uma metanálise com diferença média de −6,18 eventos/h, porém com baixa certeza e elevada heterogeneidade, e por resultados divergentes entre ensaios (Wang et al., 2020; Silva et al., 2020).',
    'Sob essas hipóteses, o cálculo em R com distribuição t não central e duas covariáveis indicou 26 participantes avaliáveis por grupo. Aplicando a previsão operacional de 20% de perdas, pela divisão de 26 por 0,80 e arredondamento para cima, definiu-se o recrutamento de 66 participantes, com meta de 33 por grupo. Essa previsão não garante a retenção de 26 participantes em cada braço. Foram realizadas análises de sensibilidade para a variabilidade, a correlação, a diferença-alvo e as perdas; a meta adotada pressupõe que esses parâmetros permaneçam próximos dos valores planejados. A retenção e as razões das perdas serão acompanhadas, preservando o plano de análise por intenção de tratar e as análises de sensibilidade para dados faltantes já especificadas.'
]
RANDOMIZATION = [
    'A sequência de alocação será gerada em R por pesquisador independente, sem participação no recrutamento, na aplicação das intervenções ou na avaliação dos desfechos, na razão 1:1, com estratificação por sexo e gravidade basal da AOS (leve: IAH ≥ 5 e < 15 eventos/h; moderada: IAH ≥ 15 e < 30 eventos/h), utilizando blocos permutados de tamanhos variáveis. Os tamanhos dos blocos, o código, a versão do R e a semente serão definidos e arquivados antes do recrutamento, sob acesso restrito. A sequência será exportada como valores fixos para uma planilha Excel, sem fórmulas aleatórias, e mantida pelo responsável independente em arquivo criptografado por senha para abertura, em armazenamento institucional com acesso restrito. A estratificação e os blocos favorecem o equilíbrio entre os grupos; a variação dos tamanhos e sua não divulgação à equipe de recrutamento reduzem a previsibilidade (Schulz e Grimes, 2002b; Hopewell et al., 2025).',
    'Após assinatura do termo de consentimento, confirmação da elegibilidade e conclusão e registro das avaliações basais, a equipe solicitará a alocação mediante identificador único e estrato do participante, por canal institucional registrado. O responsável independente utilizará a próxima posição disponível no respectivo estrato e comunicará exclusivamente à médica a atribuição daquele participante, sem compartilhar a lista completa ou as alocações futuras. O registro restrito em Excel conterá identificador, estrato, posição da lista, data e hora da solicitação e da liberação, grupo e responsável; serão preservados os registros das solicitações, versões datadas e cópias de segurança. Correções serão documentadas sem apagar o histórico; atribuições não serão trocadas nem reutilizadas. Os avaliadores utilizarão um banco separado, sem acesso à sequência ou à chave de alocação. A centralização por terceiro independente protege o sigilo no momento da inclusão e reduz o risco de viés de seleção; o cegamento de participantes e avaliadores seguirá os procedimentos já descritos (Schulz e Grimes, 2002a; Hopewell et al., 2025).'
]
REFERENCES = [
    'Vickers AJ, Altman DG. Statistics notes: Analysing controlled trials with baseline and follow up measurements. BMJ. 2001;323:1123–1124. doi:10.1136/bmj.323.7321.1123.',
    'Borm GF, Fransen J, Lemmens WA. A simple sample size formula for analysis of covariance in randomized clinical trials. J Clin Epidemiol. 2007;60:1234–1238. doi:10.1016/j.jclinepi.2007.02.006.',
    'Puhan MA, Suarez A, Lo Cascio C, Zahn A, Heitz M, Braendli O. Didgeridoo playing as alternative treatment for obstructive sleep apnoea syndrome: randomised controlled trial. BMJ. 2006;332:266–270. doi:10.1136/bmj.38705.470590.55.',
    'Silva MVFP, et al. Effects of acupuncture on obstructive sleep apnea severity, blood pressure control and quality of life in patients with hypertension: A randomized controlled trial. J Sleep Res. 2020;29:e12954. doi:10.1111/jsr.12954.',
    'Schulz KF, Grimes DA. Allocation concealment in randomised trials: defending against deciphering. Lancet. 2002a;359:614–618. doi:10.1016/S0140-6736(02)07750-4.',
    'Schulz KF, Grimes DA. Generation of allocation sequences in randomised trials: chance, not choice. Lancet. 2002b;359:515–519. doi:10.1016/S0140-6736(02)07683-3.',
    'Hopewell S, Chan AW, Collins GS, et al. CONSORT 2025 explanation and elaboration: updated guideline for reporting randomised trials. BMJ. 2025;389:e081124. doi:10.1136/bmj-2024-081124.'
]

def plain(p):
    return ''.join(p.xpath('.//w:t/text()', namespaces=NS))

def yellow(run):
    rp = run.find(Q('rPr'))
    if rp is None:
        rp = ET.Element(Q('rPr'))
        run.insert(0, rp)
    hi = rp.find(Q('highlight'))
    if hi is None:
        hi = ET.SubElement(rp, Q('highlight'))
    hi.set(Q('val'), 'yellow')

def paragraph(text, template, heading=False):
    p = ET.Element(Q('p'))
    props = deepcopy(template.find(Q('pPr')))
    if props is not None:
        if heading:
            for name in ['numPr', 'jc', 'ind']:
                for el in props.findall(Q(name)):
                    props.remove(el)
            ET.SubElement(props, Q('jc'), {Q('val'): 'left'})
            if props.find(Q('keepNext')) is None:
                ET.SubElement(props, Q('keepNext'))
        p.append(props)
    run = ET.SubElement(p, Q('r'))
    rp = ET.SubElement(run, Q('rPr'))
    ET.SubElement(rp, Q('rFonts'), {Q('ascii'): 'Arial', Q('hAnsi'): 'Arial'})
    if heading:
        ET.SubElement(rp, Q('b'))
    ET.SubElement(rp, Q('sz'), {Q('val'): '24'})
    yellow(run)
    ET.SubElement(run, Q('t')).text = text
    return p

def update_total(p, phrase):
    """Split a text run around 60; preserve its tab/other content once only."""
    found = [t for t in p.findall('.//w:t', NS) if phrase in (t.text or '')]
    assert len(found) == 1, phrase
    t = found[0]
    run = t.getparent()
    assert run.tag == Q('r') and len(run.findall(Q('t'))) == 1
    index = t.text.index(phrase)
    before, after = t.text[:index], t.text[index+2:]
    t.text = before
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    parent, position = run.getparent(), run.getparent().index(run)
    for offset, (text, mark) in enumerate([('66', True), (after, False)], 1):
        new_run = ET.Element(Q('r'))
        rp = run.find(Q('rPr'))
        if rp is not None:
            new_run.append(deepcopy(rp))
        if mark:
            yellow(new_run)
        nt = ET.SubElement(new_run, Q('t'))
        nt.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        nt.text = text
        parent.insert(position + offset, new_run)

def compatible_comments(data):
    """Flatten a Google Docs content control in comments; preserve comment text."""
    root = ET.fromstring(data)
    before = list(root.itertext())
    for el in reversed(list(root.iter(Q('sdt')))):
        content = el.find(Q('sdtContent'))
        if content is None:
            continue
        parent, index = el.getparent(), el.getparent().index(el)
        for child in list(content):
            parent.insert(index, child)
            index += 1
        parent.remove(el)
    assert list(root.itertext()) == before
    return ET.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)

with ZipFile(SOURCE) as zin:
    tree = ET.fromstring(zin.read('word/document.xml'))
    body = tree.find(Q('body'))
    ps = body.findall(Q('p'))
    changed_indices = {24, 32, 71, 79, 80}
    untouched = {i: ET.tostring(p) for i, p in enumerate(ps) if i not in changed_indices}
    original_text = {i: plain(p) for i,p in enumerate(ps)}
    assert plain(ps[71]).startswith('A sequência de randomização será gerada')
    assert plain(ps[80]).startswith('O tamanho amostral foi determinado')
    update_total(ps[24], '60 adultos')
    update_total(ps[32], '60 adults')
    update_total(ps[79], '60 indivíduos')
    body.remove(ps[71])
    inserted = []
    location = body.index(ps[78])
    for offset, p in enumerate([paragraph('Randomização e sigilo da alocação', ps[69], True)] +
                               [paragraph(s, ps[70]) for s in RANDOMIZATION]):
        body.insert(location+offset, p)
        inserted.append(p)
    location = body.index(ps[80])
    body.remove(ps[80])
    for offset, p in enumerate([paragraph('Cálculo amostral', ps[69], True)] +
                               [paragraph(s, ps[70]) for s in SAMPLE]):
        body.insert(location+offset, p)
        inserted.append(p)
    location = body.index(ps[295])
    # Freire2007 and Wang2020 already occur in the original bibliography.
    for offset, ref in enumerate(REFERENCES):
        doi = ref.split('doi:')[1].rstrip('.')
        assert doi.lower() not in '\n'.join(original_text.values()).lower(), doi
        p = paragraph(ref, ps[225])
        body.insert(location+offset, p)
        inserted.append(p)
    for i, value in untouched.items():
        assert ET.tostring(ps[i]) == value, i
    for i in [24,32,79]:
        assert plain(ps[i]) == original_text[i].replace('60 ', '66 ', 1)
    for p in inserted:
        assert all(r.find('w:rPr/w:highlight',NS).get(Q('val'))=='yellow'
                   for r in p.findall(Q('r')))
    all_text = '\n'.join(plain(p) for p in body.findall(Q('p')))
    for obsolete in ['60 adultos', '60 adults', '60 indivíduos',
                     'serão randomizados 30 participantes por grupo',
                     'seriam necessários cerca de 25 participantes']:
        assert obsolete not in all_text, obsolete
    assert all_text.count('Randomização e sigilo da alocação') == 1
    assert all_text.count('Cálculo amostral') == 1
    xml = ET.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
    with ZipFile(TARGET, 'w') as zout:
        for member in zin.infolist():
            data = xml if member.filename=='word/document.xml' else zin.read(member.filename)
            if member.filename == 'word/comments.xml':
                data = compatible_comments(data)
            zout.writestr(member, data)
with ZipFile(SOURCE) as a, ZipFile(TARGET) as b:
    assert b.testzip() is None
    assert a.namelist() == b.namelist()
    assert all(a.read(n)==b.read(n) for n in a.namelist()
               if n not in ['word/document.xml','word/comments.xml'])
    assert list(ET.fromstring(a.read('word/comments.xml')).itertext()) == list(ET.fromstring(b.read('word/comments.xml')).itertext())

client = Document()
style = client.styles['Normal']
style.font.name, style.font.size = 'Arial', Pt(11)
style.paragraph_format.space_after = Pt(10)
for chunk in (ROOT/'docs/texto_cliente_riscos.md').read_text().strip().split('\n\n'):
    if chunk.startswith('# '):
        client.add_heading(chunk[2:], 1)
    else:
        client.add_paragraph(chunk)
client.save(ROOT/'docs/texto_cliente_riscos.docx')
print(TARGET)
print('Validated: recruitment totals, unique sections, yellow highlights, original paragraphs and comment text.')
